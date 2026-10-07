#!/usr/bin/env python3
"""
Импорт журнала заказов из Excel «История заказов Эндотека.xlsx».

- События → order_journal_events (идемпотентно по import_key)
- Фото/PDF → скачивание в бакет order-attachments
- DOC/XLSX/HEIC/ZIP → ссылка (kind=url), если файл ещё доступен в RO App
- Задачи → tasks + события task_created / task_completed
- Состав работ/списания — только в журнал (остатки не трогаем)

  VITE_SUPABASE_URL=… VITE_SUPABASE_SERVICE_ROLE_KEY=… \
    python3 import/src/import-roapp-order-history.py \
      --in "/path/История заказов Эндотека.xlsx"

  --dry-run   только разбор
  --resume    продолжить с state-файла (по умолчанию вкл.)
  --no-files  не качать файлы (только события)
"""
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import os
import re
import sys
import time
import uuid
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from pathlib import Path
from threading import Lock

try:
    import openpyxl
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "--user", "openpyxl", "-q"])
    import openpyxl

MSK = timezone(timedelta(hours=3))
STORAGE_PHOTO = {".jpg", ".jpeg", ".png", ".webp", ".pdf"}
URL_ONLY = {".doc", ".docx", ".xlsx", ".xls", ".heic", ".zip", ".gif"}
DELETED_MARKERS = {"файл удалён из ro app", "файл удален из ro app", ""}

EVENT_MAP = {
    "заказ создан": "comment",
    "смена статуса": "status_changed",
    "файл прикреплён": "attachment",
    "файл прикреплен": "attachment",
    "комментарий": "comment",
    "задача создана": "task_created",
    "задача выполнена": "task_completed",
    "задача удалена": "task_deleted",
    "работа добавлена": "comment",
    "работа удалена": "comment",
    "услуга добавлена": "comment",
    "услуга удалена": "comment",
    "товар списан со склада": "comment",
    "товар возвращён на склад": "comment",
    "товар возвращен на склад": "comment",
    "товар добавлен": "comment",
    "товар удалён": "comment",
    "товар удален": "comment",
    "изменена сумма": "comment",
    "изменён клиент": "comment",
    "изменен клиент": "comment",
    "изменён срок": "comment",
    "изменен срок": "comment",
    "смена менеджера": "comment",
    "смена исполнителя": "comment",
    "заказ закрыт по балансу": "comment",
}

LABELS = {
    "заказ создан": "Заказ создан",
    "смена статуса": "Статус",
    "комментарий": "Комментарий",
    "работа добавлена": "Работа",
    "работа удалена": "Работа удалена",
    "услуга добавлена": "Услуга",
    "услуга удалена": "Услуга удалена",
    "товар списан со склада": "Списание",
    "товар возвращён на склад": "Возврат на склад",
    "товар возвращен на склад": "Возврат на склад",
    "товар добавлен": "Товар",
    "товар удалён": "Товар удалён",
    "товар удален": "Товар удалён",
    "изменена сумма": "Сумма",
    "изменён клиент": "Клиент",
    "изменен клиент": "Клиент",
    "изменён срок": "Срок",
    "изменен срок": "Срок",
    "смена менеджера": "Менеджер",
    "смена исполнителя": "Исполнитель",
    "заказ закрыт по балансу": "Закрытие",
}


def load_env():
    root = Path(__file__).resolve().parents[2]
    env_path = root / ".env"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            if "=" in line and not line.strip().startswith("#"):
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip())


def text(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\xa0", " ").strip()


def make_key(number: str, dt: datetime, event: str, extra: str) -> str:
    stamp = dt.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%S")
    raw = f"{number}|{stamp}|{event}|{extra}"
    digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]
    return f"roapp-hist:{number}:{stamp}:{digest}"


class Supabase:
    def __init__(self, url: str, key: str):
        self.url = url.rstrip("/")
        self.key = key
        self.base_headers = {
            "apikey": key,
            "Authorization": f"Bearer {key}",
            "Accept": "application/json",
        }

    def request(self, method: str, path: str, data=None, extra=None, timeout=120):
        headers = dict(self.base_headers)
        body = None
        if data is not None:
            headers["Content-Type"] = "application/json"
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        if extra:
            headers.update(extra)
        req = urllib.request.Request(self.url + path, data=body, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                raw = resp.read()
                return json.loads(raw) if raw else None
        except urllib.error.HTTPError as e:
            err = e.read()[:500].decode("utf-8", "replace")
            raise RuntimeError(f"{method} {path} → {e.code}: {err}") from e

    def get(self, path: str):
        return self.request("GET", path)

    def post(self, path: str, data, prefer="return=representation"):
        return self.request("POST", path, data, {"Prefer": prefer})

    def patch(self, path: str, data):
        return self.request("PATCH", path, data, {"Prefer": "return=minimal"})

    def upload(self, bucket: str, path: str, blob: bytes, content_type: str):
        headers = {
            "apikey": self.key,
            "Authorization": f"Bearer {self.key}",
            "Content-Type": content_type,
            "x-upsert": "true",
        }
        encoded = urllib.parse.quote(path, safe="/")
        req = urllib.request.Request(
            f"{self.url}/storage/v1/object/{bucket}/{encoded}",
            data=blob,
            headers=headers,
            method="POST",
        )
        try:
            with urllib.request.urlopen(req, timeout=180) as resp:
                resp.read()
        except urllib.error.HTTPError as e:
            err = e.read()[:400].decode("utf-8", "replace")
            raise RuntimeError(f"upload {path} → {e.code}: {err}") from e


def download(url: str) -> tuple[bytes, str]:
    req = urllib.request.Request(url, headers={"User-Agent": "endoteka-history-import/1.0"})
    with urllib.request.urlopen(req, timeout=180) as resp:
        blob = resp.read()
        ctype = resp.headers.get_content_type() or "application/octet-stream"
        return blob, ctype


def load_orders(sb: Supabase) -> dict[str, str]:
    out: dict[str, str] = {}
    for offset in range(0, 20000, 1000):
        chunk = sb.get(f"/rest/v1/orders?select=id,number&order=number&offset={offset}&limit=1000") or []
        for row in chunk:
            out[row["number"]] = row["id"]
        if len(chunk) < 1000:
            break
    return out


def load_seen_keys(sb: Supabase) -> set[str]:
    """Import keys уже в журнале (roapp-hist / roapp-history)."""
    seen: set[str] = set()
    offset = 0
    while True:
        rows = (
            sb.get(
                "/rest/v1/order_journal_events"
                f"?select=payload&offset={offset}&limit=1000"
                "&or=(payload->>source.eq.roapp-history,payload->>source.eq.roapp-hist,payload->>source.eq.roapp)"
            )
            or []
        )
        for row in rows:
            payload = row.get("payload") or {}
            if isinstance(payload, dict) and payload.get("import_key"):
                seen.add(payload["import_key"])
        if len(rows) < 1000:
            break
        offset += 1000
    return seen


def load_existing_files(sb: Supabase) -> set[str]:
    """order_id|file_name|created_at — чтобы не дублировать уже залитые файлы (A1572)."""
    out: set[str] = set()
    offset = 0
    while True:
        rows = (
            sb.get(
                "/rest/v1/order_attachments"
                f"?select=order_id,file_name,caption,created_at&offset={offset}&limit=1000"
            )
            or []
        )
        for row in rows:
            name = text(row.get("file_name") or row.get("caption") or "")
            created = text(row.get("created_at") or "")[:19]
            out.add(f"{row['order_id']}|{name}|{created}")
        if len(rows) < 1000:
            break
        offset += 1000
    return out


def parse_rows(xlsx: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    rows: list[dict] = []
    for sheet_name in wb.sheetnames:
        sh = wb[sheet_name]
        for i, row in enumerate(sh.iter_rows(values_only=True), 1):
            if i == 1:
                continue
            if not row or not row[0]:
                continue
            dt = row[1]
            if not isinstance(dt, datetime):
                continue
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=MSK)
            number = text(row[0])
            event = text(row[2])
            staff = text(row[3])
            status = text(row[4])
            desc = text(row[5])
            fname = text(row[6])
            flink = text(row[7])
            extra = fname or status or desc or event
            key = make_key(number, dt, event, extra)
            rows.append(
                {
                    "sheet": sheet_name,
                    "number": number,
                    "dt": dt,
                    "created": dt.astimezone(timezone.utc).isoformat().replace("+00:00", "Z"),
                    "event": event,
                    "event_norm": event.lower(),
                    "staff": staff,
                    "status": status,
                    "desc": desc,
                    "fname": fname,
                    "flink": flink,
                    "key": key,
                }
            )
    rows.sort(key=lambda r: (r["number"], r["dt"], r["key"]))
    return rows


def summary_for(row: dict) -> str:
    en = row["event_norm"]
    label = LABELS.get(en, row["event"])
    if en == "заказ создан":
        body = row["desc"] or "Заказ создан"
        return f"Заказ создан. {body}"[:4000]
    if en == "смена статуса":
        return (row["status"] or "Смена статуса")[:4000]
    if en == "комментарий":
        return (row["desc"] or "Комментарий")[:4000]
    if en.startswith("задача"):
        title = re.sub(r"\s*\(задача\s*#?\d+\)\s*$", "", row["desc"]).strip() or row["desc"]
        if en == "задача создана":
            return f"Создана задача: {title}"[:4000]
        if en == "задача выполнена":
            return f"Задача выполнена: {title}"[:4000]
        return f"Задача удалена: {title}"[:4000]
    if row["desc"]:
        return f"{label}: {row['desc']}"[:4000]
    return label[:4000]


def resolve_mime(fname: str, ctype: str) -> str:
    ctype = (ctype or "").split(";")[0].strip().lower()
    if ctype == "image/jpg":
        ctype = "image/jpeg"
    if ctype in {"image/jpeg", "image/png", "image/webp", "application/pdf"}:
        return ctype
    guess = mimetypes.guess_type(fname)[0] or ""
    if guess == "image/jpg":
        guess = "image/jpeg"
    if guess in {"image/jpeg", "image/png", "image/webp", "application/pdf"}:
        return guess
    ext = Path(fname).suffix.lower()
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".pdf": "application/pdf",
    }.get(ext, ctype or "application/octet-stream")


def post_url_attachment(sb: Supabase, order_id: str, row: dict, flink: str, fname: str) -> dict:
    att = sb.post(
        "/rest/v1/order_attachments",
        {
            "order_id": order_id,
            "kind": "url",
            "url": flink,
            "caption": fname,
            "file_name": fname,
            "created_at": row["created"],
        },
    )
    return att[0]


def import_file(sb: Supabase, order_id: str, row: dict, dry: bool) -> tuple[str, dict | None]:
    """Returns (mode, attachment_row_or_none). mode: storage|url|skip|deleted|error"""
    flink = row["flink"]
    fname = row["fname"] or "file"
    if flink.lower() in DELETED_MARKERS or not flink.startswith("http"):
        return "deleted", None
    ext = Path(fname).suffix.lower()
    if ext in URL_ONLY or ext not in STORAGE_PHOTO:
        if dry:
            return "url", None
        return "url", post_url_attachment(sb, order_id, row, flink, fname)
    if dry:
        return "storage", None
    try:
        blob, ctype = download(flink)
    except Exception as e:  # noqa: BLE001
        # 403/404 — сохраняем как ссылку, чтобы запись не потерялась
        print(f"WARN download {row['number']} {fname}: {e}", file=sys.stderr)
        return "url", post_url_attachment(sb, order_id, row, flink, fname)
    mime = resolve_mime(fname, ctype)
    if mime not in {"image/jpeg", "image/png", "image/webp", "application/pdf"}:
        return "url", post_url_attachment(sb, order_id, row, flink, fname)
    kind = "pdf" if mime == "application/pdf" else "photo"
    storage_ext = {".jpeg": ".jpg"}.get(ext, ext) or (".pdf" if kind == "pdf" else ".jpg")
    storage_path = f"{order_id}/{uuid.uuid4()}{storage_ext}"
    try:
        sb.upload("order-attachments", storage_path, blob, mime)
    except Exception as e:  # noqa: BLE001
        print(f"WARN upload {row['number']} {fname}: {e}", file=sys.stderr)
        return "url", post_url_attachment(sb, order_id, row, flink, fname)
    att = sb.post(
        "/rest/v1/order_attachments",
        {
            "order_id": order_id,
            "kind": kind,
            "file_path": storage_path,
            "file_name": fname,
            "mime_type": mime,
            "file_size": len(blob),
            "caption": "",
            "created_at": row["created"],
        },
    )
    return "storage", att[0]


def ensure_task(sb: Supabase, order_id: str, title: str, desc: str, created: str, cache: dict) -> str:
    cache_key = f"{order_id}:{title}"
    if cache_key in cache:
        return cache[cache_key]
    q = urllib.parse.quote(title)
    existing = sb.get(f"/rest/v1/tasks?order_id=eq.{order_id}&title=eq.{q}&select=id&limit=1") or []
    if existing:
        cache[cache_key] = existing[0]["id"]
        return cache[cache_key]
    created_task = sb.post(
        "/rest/v1/tasks",
        {
            "title": title[:200],
            "body": desc[:2000],
            "order_id": order_id,
            "priority": "normal",
            "completed": False,
            "created_at": created,
        },
    )
    cache[cache_key] = created_task[0]["id"]
    return cache[cache_key]


def main():
    load_env()
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="infile", required=True)
    parser.add_argument("--out", default="import/reports/roapp-order-history")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-files", action="store_true")
    parser.add_argument("--workers", type=int, default=10)
    parser.add_argument("--resume", action="store_true", default=True)
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args()

    url = (os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL") or "").strip()
    key = (
        os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
        or os.environ.get("VITE_SUPABASE_SERVICE_ROLE_KEY")
        or ""
    ).strip()
    if not url or not key:
        raise SystemExit("Нужны VITE_SUPABASE_URL и VITE_SUPABASE_SERVICE_ROLE_KEY")

    xlsx = Path(args.infile).expanduser().resolve()
    if not xlsx.exists():
        raise SystemExit(f"Нет файла: {xlsx}")

    out_dir = Path(args.out)
    if not out_dir.is_absolute():
        out_dir = Path(__file__).resolve().parents[2] / out_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    state_path = out_dir / "state.json"
    report_path = out_dir / "import-log.csv"

    sb = Supabase(url, key)
    print(f"CRM: {url}")
    print(f"Excel: {xlsx}")

    print("Чтение Excel…")
    rows = parse_rows(xlsx)
    if args.limit:
        rows = rows[: args.limit]
    print(f"Строк истории: {len(rows)}")

    print("Заказы CRM…")
    orders = load_orders(sb)
    print(f"Заказов в CRM: {len(orders)}")

    seen: set[str] = set()
    if args.resume and state_path.exists():
        state = json.loads(state_path.read_text(encoding="utf-8"))
        seen = set(state.get("done_keys", []))
        print(f"Resume: уже обработано {len(seen)}")
    else:
        print("Ключи из журнала…")
        seen = load_seen_keys(sb)
        print(f"Уже в БД: {len(seen)}")

    print("Уже загруженные файлы…")
    existing_files = load_existing_files(sb)
    print(f"Вложений в CRM: {len(existing_files)}")

    stats = {
        "imported": 0,
        "skipped": 0,
        "missing_order": 0,
        "files_storage": 0,
        "files_url": 0,
        "files_deleted": 0,
        "files_error": 0,
        "tasks": 0,
        "errors": 0,
    }
    missing_orders: set[str] = set()
    task_cache: dict[str, str] = {}
    log_lines = ["key,number,event,status,detail"]

    def save_state():
        state_path.write_text(
            json.dumps({"done_keys": sorted(seen), "stats": stats}, ensure_ascii=False),
            encoding="utf-8",
        )

    lock = Lock()
    started = time.time()
    pending = []
    for row in rows:
        if row["key"] in seen:
            stats["skipped"] += 1
            continue
        order_id = orders.get(row["number"])
        if not order_id:
            stats["missing_order"] += 1
            missing_orders.add(row["number"])
            seen.add(row["key"])
            log_lines.append(f'{row["key"]},{row["number"]},{row["event"]},missing,')
            continue
        pending.append((order_id, row))

    events = [(oid, r) for oid, r in pending if not r["event_norm"].startswith("файл")]
    files = [(oid, r) for oid, r in pending if r["event_norm"].startswith("файл")]
    print(f"К импорту: событий={len(events)}, файлов={len(files)}, workers={args.workers}")

    def process_event(order_id: str, row: dict) -> None:
        en = row["event_norm"]
        event_type = EVENT_MAP.get(en, "comment")
        payload = {
            "import_key": row["key"],
            "source": "roapp-hist",
            "actor_name": row["staff"],
            "ro_event": row["event"],
        }
        if en.startswith("задача"):
            title = re.sub(r"\s*\(задача\s*#?\d+\)\s*$", "", row["desc"]).strip() or "Задача из RO App"
            task_id = None
            if not args.dry_run:
                with lock:
                    task_id = ensure_task(sb, order_id, title, row["desc"], row["created"], task_cache)
                    if en == "задача выполнена":
                        sb.patch(
                            f"/rest/v1/tasks?id=eq.{task_id}",
                            {"completed": True, "completed_at": row["created"]},
                        )
                    stats["tasks"] += 1
            payload.update({"task_id": task_id, "title": title})
            detail = title
        else:
            if en == "смена статуса":
                parts = [p.strip() for p in row["status"].split("→")]
                payload.update(
                    {
                        "from": parts[0] if parts else "",
                        "to": parts[-1] if parts else row["status"],
                        "label": "Статус",
                    }
                )
            else:
                payload["body"] = row["desc"] or row["event"]
            detail = ""
        if not args.dry_run:
            sb.post(
                "/rest/v1/order_journal_events",
                {
                    "order_id": order_id,
                    "event_type": event_type,
                    "summary": summary_for(row),
                    "payload": payload,
                    "created_at": row["created"],
                },
                prefer="return=minimal",
            )
        with lock:
            seen.add(row["key"])
            stats["imported"] += 1
            log_lines.append(f'{row["key"]},{row["number"]},{row["event"]},ok,{detail}')

    def process_file(order_id: str, row: dict) -> None:
        file_stamp = row["created"][:19]
        file_dup_key = f"{order_id}|{row['fname']}|{file_stamp}"
        with lock:
            if file_dup_key in existing_files:
                seen.add(row["key"])
                stats["skipped"] += 1
                log_lines.append(
                    f'{row["key"]},{row["number"]},{row["event"]},skip_file,{row["fname"]}'
                )
                return
        payload = {
            "import_key": row["key"],
            "source": "roapp-hist",
            "actor_name": row["staff"],
            "ro_event": row["event"],
        }
        if args.no_files:
            mode, att = "skip", None
        else:
            mode, att = import_file(sb, order_id, row, args.dry_run)
        if mode == "deleted":
            summary = f"Файл удалён в RO App: {row['fname'] or 'без имени'}"
            payload.update({"kind": "deleted", "file_name": row["fname"]})
            if not args.dry_run:
                sb.post(
                    "/rest/v1/order_journal_events",
                    {
                        "order_id": order_id,
                        "event_type": "comment",
                        "summary": summary[:4000],
                        "payload": payload,
                        "created_at": row["created"],
                    },
                    prefer="return=minimal",
                )
            with lock:
                stats["files_deleted"] += 1
        elif mode in {"storage", "url"} and att:
            kind = att.get("kind") or ("url" if mode == "url" else "photo")
            summary = (
                f"Добавлена ссылка: {row['fname']}"
                if kind == "url"
                else f"Добавлен файл: {row['fname']}"
            )
            payload.update(
                {
                    "attachment_id": att["id"],
                    "kind": kind,
                    "file_name": row["fname"],
                    "mime_type": att.get("mime_type"),
                    "url": att.get("url"),
                }
            )
            if not args.dry_run:
                sb.post(
                    "/rest/v1/order_journal_events",
                    {
                        "order_id": order_id,
                        "event_type": "attachment",
                        "summary": summary[:4000],
                        "payload": payload,
                        "created_at": row["created"],
                    },
                    prefer="return=minimal",
                )
            with lock:
                if mode == "storage":
                    stats["files_storage"] += 1
                else:
                    stats["files_url"] += 1
                existing_files.add(file_dup_key)
        elif mode in {"storage", "url"} and args.dry_run:
            with lock:
                if mode == "storage":
                    stats["files_storage"] += 1
                else:
                    stats["files_url"] += 1
        with lock:
            seen.add(row["key"])
            stats["imported"] += 1
            log_lines.append(f'{row["key"]},{row["number"]},{row["event"]},{mode},{row["fname"]}')

    def run_pool(items, worker, label: str, workers: int):
        if not items:
            return
        done = 0
        total = len(items)
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            futures = {pool.submit(worker, oid, row): (oid, row) for oid, row in items}
            for fut in as_completed(futures):
                oid, row = futures[fut]
                try:
                    fut.result()
                except Exception as e:  # noqa: BLE001
                    with lock:
                        stats["errors"] += 1
                        if row["event_norm"].startswith("файл"):
                            stats["files_error"] += 1
                        log_lines.append(
                            f'{row["key"]},{row["number"]},{row["event"]},error,"{str(e).replace(chr(34), chr(39))[:200]}"'
                        )
                    print(f"ERR {row['number']} {row['event']}: {e}", file=sys.stderr)
                done += 1
                if done % 50 == 0 or done == total:
                    with lock:
                        save_state()
                        report_path.write_text("\n".join(log_lines) + "\n", encoding="utf-8")
                        snap = dict(stats)
                    elapsed = max(time.time() - started, 1)
                    rate = done / elapsed
                    eta = (total - done) / rate if rate else 0
                    print(
                        f"[{label} {done}/{total}] imported={snap['imported']} "
                        f"files={snap['files_storage']}+url={snap['files_url']} "
                        f"del={snap['files_deleted']} err={snap['errors']} "
                        f"{rate:.1f}/s eta={eta/60:.0f}m"
                    )

    print("Фаза 1/2: события без файлов…")
    run_pool(events, process_event, "events", min(8, args.workers * 2))
    print("Фаза 2/2: файлы…")
    run_pool(files, process_file, "files", args.workers)

    save_state()
    report_path.write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    missing_path = out_dir / "missing-orders.txt"
    missing_path.write_text("\n".join(sorted(missing_orders)) + ("\n" if missing_orders else ""), encoding="utf-8")

    print("DONE")
    for k, v in stats.items():
        print(f"  {k}={v}")
    print(f"  missing_orders={len(missing_orders)} → {missing_path}")
    print(f"  report={report_path}")
    print(f"  state={state_path}")


if __name__ == "__main__":
    main()
